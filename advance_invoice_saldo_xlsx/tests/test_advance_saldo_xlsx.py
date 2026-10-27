# Copyright 2026 Data Dance s.r.o.
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl-3.0).

import io

from odoo import Command
from odoo.addons.account.tests.common import AccountTestInvoicingCommon
from odoo.tests import tagged

from openpyxl import load_workbook


@tagged("post_install", "-at_install")
class TestAdvanceSaldoXlsx(AccountTestInvoicingCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env.user.group_ids |= (
            cls.env.ref("purchase.group_purchase_manager")
            + cls.env.ref("sales_team.group_sale_manager")
            + cls.env.ref("account.group_account_user")
        )
        company = cls.env.company
        Account = cls.env["account.account"]
        company.advance_paid_clearing_account_id = Account.create({
            "code": "314001", "name": "Paid advances clearing",
            "account_type": "liability_payable", "reconcile": True,
        })
        company.advance_paid_account_id = Account.create({
            "code": "314000", "name": "Paid advances",
            "account_type": "asset_current",
        })
        company.advance_purchase_journal_id = cls.env["account.journal"].create(
            {"name": "PDADV", "code": "PDADV", "type": "purchase"}
        )
        cls.service = cls.env["product.product"].create({
            "name": "Advance service", "type": "service",
        })

    def _paid_purchase_advance(self, price=500.0):
        order = self.env["purchase.order"].create({
            "partner_id": self.partner_a.id,
            "is_advance_invoice": True,
            "order_line": [
                Command.create({
                    "product_id": self.service.id,
                    "product_qty": 1,
                    "price_unit": price,
                    "tax_ids": [Command.clear()],
                })
            ],
        })
        bank_journal = self.company_data["default_journal_bank"]
        method_line = bank_journal.outbound_payment_method_line_ids[:1]
        if not method_line.payment_account_id:
            method_line.payment_account_id = self.env["account.account"].create({
                "code": "OUTPAY", "name": "Outstanding Payments",
                "account_type": "asset_current", "reconcile": True,
            })
        wizard = self.env["purchase.advance.payment.wizard"].with_context(
            active_model="purchase.order", active_id=order.id
        ).create({
            "journal_id": bank_journal.id,
            "payment_method_line_id": method_line.id,
        })
        wizard.action_create_payment()
        return order

    def _render(self, include_settled=False):
        wizard = self.env["advance.saldo.xlsx.wizard"].create({
            "include_settled": include_settled,
        })
        report = self.env.ref(
            "advance_invoice_saldo_xlsx.report_advance_saldo_xlsx"
        )
        content, report_type = self.env["ir.actions.report"]._render_xlsx(
            report.report_name, wizard.ids, data=None
        )
        self.assertEqual(report_type, "xlsx")
        return load_workbook(io.BytesIO(content), read_only=True)

    def test_open_purchase_advance_listed(self):
        order = self._paid_purchase_advance(price=500.0)
        workbook = self._render()
        sheet = workbook["Received advances"]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertGreaterEqual(len(rows), 2)
        row = next(r for r in rows[1:] if r[0] == order.name)
        self.assertEqual(row[4], 500.0)   # total
        self.assertEqual(row[5], 500.0)   # paid
        self.assertEqual(row[9], 500.0)   # open = paid - deducted
        self.assertIn("Issued advances", workbook.sheetnames)

    def test_unpaid_advance_listed_settled_hidden(self):
        order = self.env["purchase.order"].create({
            "partner_id": self.partner_a.id,
            "is_advance_invoice": True,
            "order_line": [
                Command.create({
                    "product_id": self.service.id,
                    "product_qty": 1,
                    "price_unit": 100.0,
                    "tax_ids": [Command.clear()],
                })
            ],
        })
        workbook = self._render()
        sheet = workbook["Received advances"]
        names = [r[0] for r in sheet.iter_rows(values_only=True)][1:]
        # unpaid advance still shows (payment status none)
        self.assertIn(order.name, names)

    def test_a_draft_bill_deducts_nothing_yet(self):
        """Both sheets count posted deductions only; the received side used
        to count a deduction on a bill still in draft."""
        order = self._paid_purchase_advance(price=500.0)
        bill = self.env["account.move"].create({
            "move_type": "in_invoice",
            "partner_id": self.partner_a.id,
            "invoice_date": "2026-07-15",
            "invoice_line_ids": [Command.create({
                "product_id": self.service.id, "quantity": 1,
                "price_unit": 1000.0, "tax_ids": [Command.clear()],
            })],
        })
        self.env["account.move.link.purchase.advance.wizard"].with_context(
            active_model="account.move", active_id=bill.id,
        ).create({"advance_ids": [Command.set(order.ids)]}).action_link()
        row = self._row("Received advances", order)
        self.assertEqual(row[8], 0.0)    # deducted
        self.assertEqual(row[9], 500.0)  # open
        bill.action_post()
        row = self._row("Received advances", order, include_settled=True)
        self.assertEqual(row[8], 500.0)
        self.assertEqual(row[9], 0.0)

    def test_a_standalone_issued_advance_shows_its_deduction(self):
        """Deducted with Link Advance Invoices, it used to stay open for ever."""
        company = self.env.company
        company.advance_received_account_id = self.env["account.account"].create({
            "code": "324001", "name": "Advance clearing",
            "account_type": "asset_receivable", "reconcile": True,
        })
        advance = self.env["sale.order"].create({
            "partner_id": self.partner_a.id,
            "is_advance_invoice": True,
            "order_line": [Command.create({
                "product_id": self.service.id, "product_uom_qty": 1,
                "price_unit": 300.0, "tax_ids": [Command.clear()],
            })],
        })
        advance.action_confirm()
        self.env["payment.provider"].create({
            "name": "Offline", "code": "none", "company_id": company.id,
        })
        bank = self.company_data["default_journal_bank"]
        self.env["sale.order.manual.payment.wizard"].create({
            "sale_order_id": advance.id,
            "amount": advance.amount_total,
            "journal_id": bank.id,
            "payment_method_line_id":
                bank._get_available_payment_method_lines("inbound")[:1].id,
        }).action_create_payment()
        invoice = self.env["account.move"].create({
            "move_type": "out_invoice",
            "partner_id": self.partner_a.id,
            "invoice_line_ids": [Command.create({
                "product_id": self.service.id, "quantity": 1,
                "price_unit": 1000.0, "tax_ids": [Command.clear()],
            })],
        })
        self.env["account.move.link.advance.invoice.wizard"].with_context(
            active_model="account.move", active_id=invoice.id,
        ).create({
            "move_id": invoice.id,
            "advance_invoice_ids": [Command.set(advance.ids)],
        }).action_link_advance_invoices()
        invoice.action_post()
        row = self._row("Issued advances", advance, include_settled=True)
        self.assertEqual(row[5], 300.0)  # paid
        self.assertEqual(row[8], 300.0)  # deducted
        self.assertEqual(row[9], 0.0)    # open

    def test_the_headings_follow_the_user_s_language(self):
        for code in ("sk_SK", "cs_CZ"):
            self.env["res.lang"]._activate_lang(code)
        report = self.env["report.advance_invoice_saldo_xlsx.saldo"]
        self.assertEqual(report.with_context(lang="sk_SK")._headers()[0], "Číslo")
        self.assertEqual(report.with_context(lang="cs_CZ")._headers()[0], "Číslo")
        self.assertEqual(report.with_context(lang="en_US")._headers()[0], "Number")

    def _row(self, sheet_name, order, include_settled=False):
        sheet = self._render(include_settled=include_settled)[sheet_name]
        return next(
            row for row in sheet.iter_rows(values_only=True) if row[0] == order.name
        )
