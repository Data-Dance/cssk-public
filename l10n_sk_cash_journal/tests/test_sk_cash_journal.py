# Copyright 2026 Data Dance s.r.o.
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

from odoo import fields
from odoo.addons.account.tests.common import AccountTestInvoicingCommon
from odoo.tests import tagged


@tagged("post_install", "-at_install")
class TestSkCashJournal(AccountTestInvoicingCommon):
    """The Slovak catalogue, the statutory grid and the DPFO figures."""

    @classmethod
    @AccountTestInvoicingCommon.setup_country("sk")
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.company.cssk_bookkeeping_regime = "de"

        def category(xmlid):
            return cls.env.ref("l10n_sk_cash_journal.%s" % xmlid)

        cls.cat_vyrobky = category("cat_p_vyrobky")
        cls.cat_zasoby = category("cat_v_zasoby")
        cls.cat_sluzby = category("cat_v_sluzby")
        cls.cat_poistne = category("cat_v_poistne_podnikatel")
        cls.cat_osobna = category("cat_vn_osobna_spotreba")
        cls.cat_priebezne = category("cat_c_priebezne")
        cls.cat_odpisy = category("cat_z_odpisy")

        cls.income_account = cls.company_data["default_account_revenue"]
        cls.expense_account = cls.company_data["default_account_expense"]
        cls.income_account.cssk_cash_category_id = cls.cat_vyrobky
        cls.expense_account.cssk_cash_category_id = cls.cat_zasoby

        cls.bank_journal = cls.company_data["default_journal_bank"]
        cls.cash_journal = cls.company_data["default_journal_cash"]
        cls.bank_journal.default_account_id.cssk_cash_category_id = \
            cls.cat_priebezne
        cls.cash_journal.default_account_id.cssk_cash_category_id = \
            cls.cat_priebezne

    # -- helpers -------------------------------------------------------

    def _invoice(self, move_type="out_invoice", amount=1000.0, account=None,
                 invoice_date="2026-02-01"):
        invoice = self.env["account.move"].create({
            "move_type": move_type,
            "partner_id": self.partner_a.id,
            "invoice_date": fields.Date.to_date(invoice_date),
            "date": fields.Date.to_date(invoice_date),
            "invoice_line_ids": [(0, 0, {
                "name": "položka",
                "quantity": 1.0,
                "price_unit": amount,
                "account_id": (account or self.income_account).id,
                "tax_ids": [(5, 0, 0)],
            })],
        })
        invoice.action_post()
        return invoice

    def _pay(self, invoice, journal=None, payment_date="2026-03-15"):
        wizard = self.env["account.payment.register"].with_context(
            active_model="account.move", active_ids=invoice.ids,
        ).create({
            "payment_date": fields.Date.to_date(payment_date),
            "journal_id": (journal or self.bank_journal).id,
        })
        return wizard._create_payments()

    def _generate(self):
        return self.env["cssk.cash.journal.line"]._cssk_regenerate(
            self.company,
            fields.Date.to_date("2026-01-01"),
            fields.Date.to_date("2026-12-31"),
        )

    # -- the default chart mapping -------------------------------------

    def _mapped(self, code):
        account = self.env["account.account"].with_company(self.company).search([
            ("company_ids", "in", self.company.id),
            ("code", "=like", code + "%"),
        ], limit=1)
        self.assertTrue(account, "account %s is missing from the chart" % code)
        return account

    def test_loading_the_chart_maps_it_without_being_asked(self):
        """The usual order is module first, chart second.

        This company's chart was loaded by the test fixture, with no call to the
        mapping at all — so if the accounts below carry categories, the chart
        hook did it. Measured on a fresh database: 89 accounts outbound and 15
        with an inbound category.
        """
        self.assertEqual(self._mapped("501").cssk_cash_category_id.code, "V1")
        self.assertEqual(self._mapped("343").cssk_cash_category_in_id.code, "PN3")
        mapped = self.env["account.account"].with_company(self.company).search_count([
            ("company_ids", "in", self.company.id),
            ("cssk_cash_category_id", "!=", False),
        ])
        self.assertGreater(mapped, 50, "the whole chart, not a handful")

    def test_the_chart_maps_itself_on_install(self):
        """An accountant should see a denník without mapping 40 accounts first."""
        self.company._cssk_map_chart_categories()
        self.assertEqual(
            self._mapped("501").cssk_cash_category_id.code, "V1")
        self.assertEqual(
            self._mapped("604").cssk_cash_category_id.code, "P1")
        self.assertEqual(
            self._mapped("602").cssk_cash_category_id.code, "P2")
        self.assertEqual(
            self._mapped("518").cssk_cash_category_id.code, "V2")
        self.assertEqual(
            self._mapped("521").cssk_cash_category_id.code, "V3")

    def test_poistne_podnikatela_is_mapped_as_a_tax_expense(self):
        """526 is where the two countries part company (§ 19 ods. 3 písm. i)."""
        self.company._cssk_map_chart_categories()
        category = self._mapped("526").cssk_cash_category_id
        self.assertEqual(category.code, "V4")
        self.assertTrue(category.taxable)

    def test_depreciation_is_mapped_to_the_non_cash_category(self):
        self.company._cssk_map_chart_categories()
        category = self._mapped("551").cssk_cash_category_id
        self.assertEqual(category.code, "Z1")
        self.assertTrue(category.non_cash)

    def test_two_way_accounts_are_mapped_in_both_directions(self):
        """343, 461 and 491 each mean a different column in each direction."""
        self.company._cssk_map_chart_categories()
        for code, out_code, in_code in (
            ("343", "VN3", "PN3"),
            ("461", "VN5", "PN2"),
            ("491", "VN1", "PN1"),
        ):
            account = self._mapped(code)
            self.assertEqual(account.cssk_cash_category_id.code, out_code, code)
            self.assertEqual(
                account.cssk_cash_category_in_id.code, in_code, code)

    def test_receivables_payables_and_advances_are_left_unmapped(self):
        """The engine follows those to the document; an advance needs judgment."""
        self.company._cssk_map_chart_categories()
        for code in ("311", "321"):
            account = self._mapped(code)
            self.assertFalse(account.cssk_cash_category_id, code)
            self.assertFalse(account.cssk_cash_category_in_id, code)

    def test_mapping_never_overwrites_a_choice_already_made(self):
        account = self._mapped("501")
        account.cssk_cash_category_id = self.env.ref(
            "l10n_sk_cash_journal.cat_v_ostatne")
        self.company._cssk_map_chart_categories()
        self.assertEqual(account.cssk_cash_category_id.code, "V9",
                         "the accountant's own decision outranks the default")
        self.company._cssk_map_chart_categories(overwrite=True)
        self.assertEqual(account.cssk_cash_category_id.code, "V1",
                         "and can be reset deliberately")

    # -- the catalogue -------------------------------------------------

    def test_the_statutory_categories_are_installed(self):
        """Opatrenie MF/27076/2007-74, § 4 — every column has a category."""
        categories = self.env["cssk.cash.category"].search([
            ("country_id", "=", self.env.ref("base.sk").id),
        ])
        codes = set(categories.mapped("code"))
        self.assertTrue(
            {"P1", "P2", "P3", "V1", "V2", "V3", "V4", "V5", "V6", "V9",
             "VN1", "VN2", "VN3", "C1", "Z1"} <= codes)

    def test_poistne_podnikatela_is_a_tax_expense_in_slovakia(self):
        """§ 19 ods. 3 písm. i) ZDP — the clearest CZ/SK difference.

        In the Czech Republic the same payment is non-deductible (§ 25 odst. 1
        písm. g), so the two catalogues must disagree here.
        """
        self.assertTrue(self.cat_poistne.taxable)
        self.assertEqual(self.cat_poistne.kind, "expense")

    def test_depreciation_is_a_non_cash_category(self):
        self.assertTrue(self.cat_odpisy.non_cash)
        self.assertTrue(self.cat_odpisy.taxable)

    # -- the grid ------------------------------------------------------

    def test_the_grid_puts_a_sale_in_its_statutory_column(self):
        invoice = self._invoice()
        self._pay(invoice)
        self._generate()
        values = self.env["report.l10n_sk_cash_journal.report_penazny_dennik"] \
            ._sk_dennik_values(
                self.company,
                fields.Date.to_date("2026-01-01"),
                fields.Date.to_date("2026-12-31"))
        self.assertEqual(len(values["lines"]), 1)
        cells = values["lines"][0]["cells"]
        self.assertAlmostEqual(cells["p_vyrobky"], 1000.0, places=2)
        self.assertAlmostEqual(cells["banka_prijem"], 1000.0, places=2)
        self.assertAlmostEqual(values["totals"]["p_vyrobky"], 1000.0, places=2)

    def test_the_running_balance_follows_the_bank(self):
        """§ 4 ods. 10 — the book must tie to the cash and bank balances."""
        sale = self._invoice(amount=1000.0)
        self._pay(sale, payment_date="2026-03-01")
        bill = self._invoice(move_type="in_invoice", amount=400.0,
                             account=self.expense_account)
        self._pay(bill, payment_date="2026-03-10")
        self._generate()
        values = self.env["report.l10n_sk_cash_journal.report_penazny_dennik"] \
            ._sk_dennik_values(
                self.company,
                fields.Date.to_date("2026-01-01"),
                fields.Date.to_date("2026-12-31"))
        self.assertAlmostEqual(values["closing_bank"], 600.0, places=2)

        # "The bank" is the bank account plus the transit accounts: without a
        # statement the money is still sitting on outstanding receipts, and with
        # one the two transit legs net out.
        money_accounts = self.bank_journal.default_account_id | self.env[
            "cssk.cash.journal.line"]._cssk_transit_accounts(self.company)
        ledger = self.env["account.move.line"].search([
            ("company_id", "=", self.company.id),
            ("parent_state", "=", "posted"),
            ("account_id", "in", money_accounts.ids),
        ])
        self.assertAlmostEqual(
            values["closing_bank"], sum(ledger.mapped("balance")), places=2,
            msg="the denník must reconcile to the bank account")

    def test_a_category_the_layout_does_not_name_still_appears(self):
        """Otherwise the columns would not add up to the money."""
        private = self.expense_account.copy({
            "code": "501900",
            "cssk_cash_category_id": self.cat_osobna.id,
        })
        bill = self._invoice(move_type="in_invoice", amount=250.0,
                             account=private)
        self._pay(bill)
        self._generate()
        values = self.env["report.l10n_sk_cash_journal.report_penazny_dennik"] \
            ._sk_dennik_values(
                self.company,
                fields.Date.to_date("2026-01-01"),
                fields.Date.to_date("2026-12-31"))
        self.assertAlmostEqual(values["totals"]["vn"], 250.0, places=2)

    # -- the tax return ------------------------------------------------

    def test_dpfo_tabulka_1_totals_come_from_the_book(self):
        sale = self._invoice(amount=1000.0)
        self._pay(sale, payment_date="2026-03-01")
        bill = self._invoice(move_type="in_invoice", amount=400.0,
                             account=self.expense_account)
        self._pay(bill, payment_date="2026-03-10")
        self._generate()

        wizard = self.env["l10n.sk.cash.dpfo"].create({
            "company_id": self.company.id,
            "date_from": fields.Date.to_date("2026-01-01"),
            "date_to": fields.Date.to_date("2026-12-31"),
        })
        wizard.action_compute()
        figures = {line.code: line.value for line in wizard.line_ids}
        self.assertAlmostEqual(figures["t1r10_prijmy"], 1000.0, places=2)
        self.assertAlmostEqual(figures["t1r10_vydavky"], 400.0, places=2)
        self.assertAlmostEqual(figures["zaklad"], 600.0, places=2)
        self.assertAlmostEqual(figures["t1r2"], 1400.0, places=2,
                               msg="both sides of § 6 ods. 1 b) carry the code")

    def test_a_non_taxable_payment_stays_out_of_the_tax_base(self):
        """An owner's withdrawal is in the book and not in the return."""
        private = self.expense_account.copy({
            "code": "501910",
            "cssk_cash_category_id": self.cat_osobna.id,
        })
        bill = self._invoice(move_type="in_invoice", amount=300.0,
                             account=private)
        self._pay(bill)
        self._generate()
        wizard = self.env["l10n.sk.cash.dpfo"].create({
            "company_id": self.company.id,
            "date_from": fields.Date.to_date("2026-01-01"),
            "date_to": fields.Date.to_date("2026-12-31"),
        })
        wizard.action_compute()
        figures = {line.code: line.value for line in wizard.line_ids}
        self.assertAlmostEqual(figures["t1r10_vydavky"], 0.0, places=2)

    def test_dpfo_tabulka_1a_balances_come_from_the_ledger(self):
        """Tabuľka 1a wants the state at both ends of the year."""
        self._invoice(amount=1000.0, invoice_date="2026-06-01")
        wizard = self.env["l10n.sk.cash.dpfo"].create({
            "company_id": self.company.id,
            "date_from": fields.Date.to_date("2026-01-01"),
            "date_to": fields.Date.to_date("2026-12-31"),
        })
        wizard.action_compute()
        rows = {line.code: line for line in wizard.line_ids
                if line.table == "t1a"}
        self.assertEqual(
            set(rows), {"r1", "r2", "r3", "r4", "r5", "r6"})
        self.assertAlmostEqual(rows["r4"].opening_value, 0.0, places=2)
        self.assertAlmostEqual(rows["r4"].closing_value, 1000.0, places=2,
                               msg="an unpaid invoice is a pohľadávka at year end")

    def test_a_refunded_sale_reduces_the_sales_column_and_tabulka_1(self):
        """The accountant's dobropis question, end to end.

        Measured before the fix: the sales column read 1200 and tabuľka 1 said
        príjmy 1000 against výdavky 200. The statute wants the gross columns,
        so a refund reduces príjmy and no expense appears.
        """
        sale = self._invoice(amount=1000.0)
        self._pay(sale, payment_date="2026-03-01")
        refund = self.env["account.move"].create({
            "move_type": "out_refund",
            "partner_id": self.partner_a.id,
            "invoice_date": fields.Date.to_date("2026-03-04"),
            "date": fields.Date.to_date("2026-03-04"),
            "invoice_line_ids": [(0, 0, {
                "name": "vrátenie",
                "quantity": 1.0,
                "price_unit": 200.0,
                "account_id": self.income_account.id,
                "tax_ids": [(5, 0, 0)],
            })],
        })
        refund.action_post()
        self._pay(refund, payment_date="2026-03-05")
        self._generate()

        values = self.env["report.l10n_sk_cash_journal.report_penazny_dennik"] \
            ._sk_dennik_values(
                self.company,
                fields.Date.to_date("2026-01-01"),
                fields.Date.to_date("2026-12-31"))
        self.assertAlmostEqual(values["totals"]["p_vyrobky"], 800.0, places=2)
        self.assertAlmostEqual(values["totals"]["banka_prijem"], 1000.0, places=2)
        self.assertAlmostEqual(values["totals"]["banka_vydaj"], 200.0, places=2)
        self.assertAlmostEqual(values["closing_bank"], 800.0, places=2)
        for key in ("v_ostatne", "vn"):
            self.assertAlmostEqual(
                values["totals"][key], 0.0, places=2,
                msg="a refunded sale is not an expense")

        wizard = self.env["l10n.sk.cash.dpfo"].create({
            "company_id": self.company.id,
            "date_from": fields.Date.to_date("2026-01-01"),
            "date_to": fields.Date.to_date("2026-12-31"),
        })
        wizard.action_compute()
        figures = {line.code: line.value for line in wizard.line_ids}
        self.assertAlmostEqual(figures["t1r10_prijmy"], 800.0, places=2)
        self.assertAlmostEqual(figures["t1r10_vydavky"], 0.0, places=2)
        self.assertAlmostEqual(figures["zaklad"], 800.0, places=2)

    def test_a_loan_shows_in_both_neovplyvnujuce_columns(self):
        """The accountant's pôžička question, in the statutory grid.

        One account, two columns: PN2 when the loan arrives, VN5 when an
        instalment leaves. Mapped one way it used to net into a single column.
        """
        loan = self.env["account.account"].create({
            "name": "Bankový úver",
            "code": "461900",
            "account_type": "liability_non_current",
            "reconcile": True,
            "cssk_cash_category_id": self.env.ref(
                "l10n_sk_cash_journal.cat_vn_uver").id,
            "cssk_cash_category_in_id": self.env.ref(
                "l10n_sk_cash_journal.cat_pn_uver").id,
        })
        bank = self.bank_journal.default_account_id
        for amount, incoming, date in ((10000.0, True, "2026-05-02"),
                                       (500.0, False, "2026-06-02")):
            move = self.env["account.move"].create({
                "journal_id": self.bank_journal.id,
                "date": fields.Date.to_date(date),
                "line_ids": [
                    (0, 0, {"account_id": bank.id, "name": "úver",
                            "debit": amount if incoming else 0.0,
                            "credit": 0.0 if incoming else amount}),
                    (0, 0, {"account_id": loan.id, "name": "úver",
                            "debit": 0.0 if incoming else amount,
                            "credit": amount if incoming else 0.0}),
                ],
            })
            move.action_post()
        self._generate()

        values = self.env["report.l10n_sk_cash_journal.report_penazny_dennik"] \
            ._sk_dennik_values(
                self.company,
                fields.Date.to_date("2026-01-01"),
                fields.Date.to_date("2026-12-31"))
        self.assertAlmostEqual(values["totals"]["pn"], 10000.0, places=2)
        self.assertAlmostEqual(values["totals"]["vn"], 500.0, places=2)
        self.assertAlmostEqual(values["closing_bank"], 9500.0, places=2)

        wizard = self.env["l10n.sk.cash.dpfo"].create({
            "company_id": self.company.id,
            "date_from": fields.Date.to_date("2026-01-01"),
            "date_to": fields.Date.to_date("2026-12-31"),
        })
        wizard.action_compute()
        figures = {line.code: line.value for line in wizard.line_ids}
        self.assertAlmostEqual(figures["t1r10_prijmy"], 0.0, places=2)
        self.assertAlmostEqual(figures["t1r10_vydavky"], 0.0, places=2,
                               msg="a loan touches neither side of the tax base")

    def test_a_flat_rate_payer_gets_tabulka_1b(self):
        """§ 6 ods. 10 with § 6 ods. 11 a) and d): zásoby and pohľadávky only."""
        self.company.cssk_bookkeeping_regime = "pausal"
        wizard = self.env["l10n.sk.cash.dpfo"].create({
            "company_id": self.company.id,
            "date_from": fields.Date.to_date("2026-01-01"),
            "date_to": fields.Date.to_date("2026-12-31"),
        })
        wizard.action_compute()
        tables = {line.table for line in wizard.line_ids}
        self.assertIn("t1b", tables)
        self.assertNotIn("t1a", tables)
        codes = {line.code for line in wizard.line_ids if line.table == "t1b"}
        self.assertEqual(codes, {"r3", "r4"})

    # -- the statutory sheet of the spreadsheet ------------------------
    def test_the_spreadsheet_opens_with_the_statutory_grid(self):
        """"Zvlášť stĺpce pre príjem a výdaj v pokladni a zvlášť príjem a
        výdaj v banke. Tak ako je to v tom pdf." Same figures as the PDF."""
        import io

        from openpyxl import load_workbook

        self._pay(self._invoice(amount=1000.0))
        self._pay(self._invoice(move_type="in_invoice", amount=200.0,
                                account=self.expense_account),
                  journal=self.cash_journal)
        self._generate()
        wizard = self.env["cssk.cash.journal.print"].create({
            "company_id": self.company.id, "output": "xlsx",
            "date_from": "2026-01-01", "date_to": "2026-12-31"})
        content, _ext = self.env["ir.actions.report"]._render_xlsx(
            "l10n_cssk_cash_journal_base.dennik_xlsx", wizard.ids, data=None)
        book = load_workbook(io.BytesIO(content), data_only=True)
        self.assertEqual(book.sheetnames[0], "Peňažný denník")
        self.assertIn("Cash journal", " ".join(book.sheetnames[1:]),
                      "the flat sheet for pivoting stays")
        sheet = book.worksheets[0]
        header = [c.value for c in sheet[1]]
        sub = [c.value for c in sheet[2]]
        for group in ("Pokladnica", "Banka", "Priebežné položky", "DPH"):
            col = header.index(group)
            self.assertEqual(sub[col:col + 2], ["príjem", "výdaj"], group)
        values = self.env["report.l10n_sk_cash_journal.report_penazny_dennik"] \
            ._sk_dennik_values(self.company, fields.Date.to_date("2026-01-01"),
                               fields.Date.to_date("2026-12-31"))
        total_row = [c.value for c in list(sheet.rows)[-1]]
        self.assertEqual(total_row[3], "Spolu")
        self.assertAlmostEqual(total_row[header.index("Banka")],
                               values["totals"]["banka_prijem"], 2)
        self.assertAlmostEqual(total_row[header.index("Pokladnica") + 1],
                               values["totals"]["pokladnica_vydaj"], 2)
        self.assertAlmostEqual(values["totals"]["banka_prijem"], 1000.0, 2)
        self.assertAlmostEqual(values["totals"]["pokladnica_vydaj"], 200.0, 2)
        self.assertAlmostEqual(total_row[-1], values["closing_bank"], 2)

